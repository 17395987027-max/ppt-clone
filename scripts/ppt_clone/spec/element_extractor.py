"""Element-level extraction — capture every visual detail from python-pptx shapes.

For each shape on a slide, produces an Element with:
  - Text: font_family, font_size_pt, bold, italic, color, alignment, line_spacing
  - Shape: fill_color, stroke_color, stroke_width, corner_radius, shape_type
  - Image: src path, crop info
  - Position: normalized (0–1) coordinates
  - Semantic role: title, body, subtitle, footer, decoration, etc.
"""

from __future__ import annotations

import re
import io
import os
from pathlib import Path
from typing import Any

from ppt_clone.spec.spec_model import (
    Element,
    ElementType,
    ImageStyle,
    Pos,
    SemanticRole,
    ShapeStyle,
    TextStyle,
)


# ── EMU conversion ───────────────────────────────────────────────────

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
PT_PER_EMU = 1.0 / EMU_PER_PT


def emu_to_inches(emu: int) -> float:
    return emu / EMU_PER_INCH


def emu_to_pt(emu: int) -> float:
    return emu * PT_PER_EMU


# ── Theme colour resolution ──────────────────────────────────────────
#
# python-pptx exposes fill/line/font colours as ColorFormat. When the colour
# is an explicit srgb value, ``.rgb`` works. When it is a *scheme* colour
# (the common case in branded decks — accent1, dk1, tx1 …), ``.rgb`` raises
# and the caller used to swallow the error, leaving the spec colour blank.
# Here we resolve scheme colours against the deck's theme palette
# (extract_theme_colors) so the spec carries real HEX values.

# OOXML slide scheme keys -> clrScheme keys (tx1/bg1 alias dk1/lt1 …)
_THEME_KEY_ALIASES = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}

# Active palette, set per-extraction by set_theme_colors().
_THEME_COLORS: dict[str, str] = {}


def set_theme_colors(theme_colors: dict | None) -> None:
    """Set the theme palette used to resolve scheme colours to HEX."""
    global _THEME_COLORS
    _THEME_COLORS = {k: v for k, v in (theme_colors or {}).items() if v}


def _scheme_key(color_format) -> str | None:
    """Return the clrScheme key ('accent1', 'dk1', …) for a scheme colour."""
    try:
        tc = color_format.theme_color
    except Exception:
        return None
    if tc is None:
        return None
    key = getattr(tc, "xml_value", None)
    return key or None


def resolve_to_hex(color_format) -> str:
    """Resolve a ColorFormat to '#RRGGBB', or '' when it cannot be resolved.

    Handles explicit srgb colours and scheme colours mapped through the
    active theme palette. Returns '' for unset/inherited colours.
    """
    if color_format is None:
        return ""
    # 1) explicit RGB value
    try:
        rgb = color_format.rgb
        if rgb is not None:
            return f"#{rgb}"
    except Exception:
        pass
    # 2) scheme colour -> theme palette
    key = _scheme_key(color_format)
    if key:
        hexv = _THEME_COLORS.get(key) or _THEME_COLORS.get(
            _THEME_KEY_ALIASES.get(key, ""), ""
        )
        if hexv:
            return hexv if hexv.startswith("#") else f"#{hexv}"
    return ""


# ── XML colour elements (with luminance/shade/tint modifiers) ─────────
#
# python-pptx's ColorFormat cannot see the lumMod/lumOff/shade/tint child
# elements DrawingML uses for tints and shades (e.g. accent1 lightened 40%).
# Reading the raw colour element lets us reproduce those exact values.

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# A few common preset colours (prstClr) — enough for practical decks.
_PRESET_COLORS = {
    "black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000",
    "blue": "0000FF", "yellow": "FFFF00", "gray": "808080", "grey": "808080",
    "orange": "FFA500", "purple": "800080", "silver": "C0C0C0",
}


def _apply_tint_shade(hexstr: str, elem) -> str:
    """Apply lumMod/lumOff/shade/tint children of a colour element to a HEX."""
    import colorsys

    try:
        r = int(hexstr[1:3], 16) / 255.0
        g = int(hexstr[3:5], 16) / 255.0
        b = int(hexstr[5:7], 16) / 255.0
    except Exception:
        return hexstr

    lum_mod = lum_off = sh = tn = None
    for ch in elem:
        tag = ch.tag.split("}")[-1]
        try:
            v = int(ch.get("val", "0")) / 100000.0
        except Exception:
            continue
        if tag == "lumMod":
            lum_mod = v
        elif tag == "lumOff":
            lum_off = v
        elif tag == "shade":
            sh = v
        elif tag == "tint":
            tn = v

    # luminance transforms operate in HSL lightness
    if lum_mod is not None or lum_off is not None:
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        if lum_mod is not None:
            l *= lum_mod
        if lum_off is not None:
            l += lum_off
        l = min(1.0, max(0.0, l))
        r, g, b = colorsys.hls_to_rgb(h, l, s)

    # shade darkens toward black, tint lightens toward white
    if sh is not None:
        r, g, b = r * sh, g * sh, b * sh
    if tn is not None:
        r = r + (1.0 - r) * tn
        g = g + (1.0 - g) * tn
        b = b + (1.0 - b) * tn

    return "#%02X%02X%02X" % (
        round(min(1, max(0, r)) * 255),
        round(min(1, max(0, g)) * 255),
        round(min(1, max(0, b)) * 255),
    )


def resolve_color_element(color_el) -> str:
    """Resolve an <a:srgbClr>/<a:schemeClr>/<a:sysClr>/<a:prstClr> to HEX.

    Applies any lumMod/lumOff/shade/tint modifiers, so tints/shades of theme
    colours come out exact instead of approximated by the base colour.
    """
    if color_el is None:
        return ""
    tag = color_el.tag.split("}")[-1]
    base = ""
    if tag == "srgbClr":
        base = "#" + (color_el.get("val") or "")
    elif tag == "schemeClr":
        key = color_el.get("val") or ""
        hexv = _THEME_COLORS.get(key) or _THEME_COLORS.get(
            _THEME_KEY_ALIASES.get(key, ""), ""
        )
        if not hexv:
            return ""
        base = hexv if hexv.startswith("#") else f"#{hexv}"
    elif tag == "sysClr":
        base = "#" + (color_el.get("lastClr") or "000000")
    elif tag == "prstClr":
        base = "#" + _PRESET_COLORS.get(color_el.get("val", ""), "")
    else:
        return ""
    if len(base) != 7 or base == "#":
        return ""
    return _apply_tint_shade(base, color_el)


def _first_color_el(solid_fill) -> any:
    """Return the colour element inside an <a:solidFill>, if any."""
    if solid_fill is None:
        return None
    for child in solid_fill:
        if child.tag.split("}")[-1] in ("srgbClr", "schemeClr", "sysClr", "prstClr"):
            return child
    return None


def color_alpha(color_el) -> float:
    """Return the opacity (0–1) declared on a colour element via <a:alpha>.

    OOXML expresses transparency as an opacity value in 1/1000 of a percent
    (``val="20000"`` → 20% opacity). Branded decks use this heavily for the
    dimming overlays and translucent bands painted on layouts, so the value
    is part of the visual spec, not a detail to approximate away.
    """
    if color_el is None:
        return 1.0
    alpha = color_el.find(f"{{{_NS_A}}}alpha")
    if alpha is None:
        return 1.0
    try:
        return min(1.0, max(0.0, int(alpha.get("val", "100000")) / 100000.0))
    except Exception:
        return 1.0


def extract_gradient(grad_fill) -> tuple[list[dict], float | None]:
    """Return (stops, angle_deg) from an <a:gradFill> element."""
    if grad_fill is None:
        return [], None
    stops: list[dict] = []
    gs_lst = grad_fill.find(f"{{{_NS_A}}}gsLst")
    if gs_lst is not None:
        for gs in gs_lst.findall(f"{{{_NS_A}}}gs"):
            try:
                pos = int(gs.get("pos", "0")) / 100000.0
            except Exception:
                pos = 0.0
            col = resolve_color_element(_first_color_el(gs))
            if col:
                stops.append({"pos": round(pos, 3), "color": col})
    angle = None
    lin = grad_fill.find(f"{{{_NS_A}}}lin")
    if lin is not None:
        try:
            angle = int(lin.get("ang", "0")) / 60000.0
        except Exception:
            angle = None
    return stops, angle


# ── Main extraction ──────────────────────────────────────────────────


def extract_element(shape, slide_width_emu: int,
                    slide_height_emu: int,
                    z_order: int,
                    spec_dir: Path) -> Element:
    """Extract complete element info from a python-pptx shape.

    Returns Element with detailed text/shape/image styles.
    """
    sw = slide_width_emu
    sh = slide_height_emu

    pos = Pos(
        x=(shape.left or 0) / sw if sw else 0,
        y=(shape.top or 0) / sh if sh else 0,
        w=(shape.width or 0) / sw if sw else 0,
        h=(shape.height or 0) / sh if sh else 0,
    )

    shape_name = getattr(shape, "name", "")

    # Determine semantic role
    role = _infer_role(shape)

    # Build element
    elem = Element(
        id=z_order,
        element_type=_get_element_type(shape),
        semantic_role=role,
        position=pos,
        z_order=z_order,
        shape_name=shape_name,
    )

    # Extract text details
    try:
        if shape.has_text_frame:
            elem.text = shape.text_frame.text
            elem.text_style = _extract_text_style(shape.text_frame)
            # Build text from all paragraphs, skipping empty leading ones
            parts = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if parts:
                elem.text = "\n".join(parts)
    except (AttributeError, ValueError):
        pass

    # Extract shape details
    elem.shape_style = _extract_shape_style(shape)

    # Extract image details
    try:
        from pptx.shapes.picture import Picture
        if isinstance(shape, Picture):
            elem.image_style = _extract_image_style(shape, spec_dir, z_order)
            if elem.image_style:
                elem.element_type = ElementType.IMAGE
                elem.semantic_role = SemanticRole.IMAGE
    except ImportError:
        pass

    # Handle group shapes
    if hasattr(shape, "shapes"):
        for ci, child in enumerate(shape.shapes):
            child_elem = extract_element(child, sw, sh, ci, spec_dir)
            if child_elem.element_type == ElementType.TEXT and not elem.text:
                elem.text = child_elem.text
                elem.text_style = child_elem.text_style
            elem.children.append(child_elem)

    # Bubble up text from children
    if not elem.text and elem.children:
        for child in elem.children:
            if child.text:
                elem.text = child.text
                elem.text_style = child.text_style
                break

    return elem


# ── Element type detection ────────────────────────────────────────────


def _get_element_type(shape) -> ElementType:
    try:
        if hasattr(shape, "has_table") and shape.has_table:
            return ElementType.TABLE
    except Exception: pass
    try:
        if hasattr(shape, "chart"):
            shape.chart
            return ElementType.CHART
    except Exception: pass
    try:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return ElementType.TEXT
    except Exception: pass
    if hasattr(shape, "shapes"):
        return ElementType.GROUP
    return ElementType.SHAPE


# ── Semantic role inference ──────────────────────────────────────────


def _infer_role(shape) -> SemanticRole:
    name = getattr(shape, "name", "").lower()
    text = shape.text_frame.text if shape.has_text_frame else ""

    if shape.has_text_frame:
        text_lower = text.lower()
        text_len = len(text.strip())
        try:
            pf = shape.placeholder_format
            if pf is not None:
                from pptx.enum.shapes import PP_PLACEHOLDER
                ph_type = pf.type
                if ph_type == PP_PLACEHOLDER.TITLE:
                    return SemanticRole.TITLE
                elif ph_type == PP_PLACEHOLDER.SUBTITLE:
                    return SemanticRole.SUBTITLE
                elif ph_type == PP_PLACEHOLDER.BODY:
                    return SemanticRole.BODY
                elif ph_type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE):
                    return SemanticRole.PAGE_NUMBER
        except (ValueError, AttributeError, Exception):
            pass

    # Name-based inference
    if any(kw in name for kw in ["title", "标题", "heading"]):
        return SemanticRole.TITLE
    if any(kw in name for kw in ["subtitle", "副标题", "sub"]):
        return SemanticRole.SUBTITLE
    if any(kw in name for kw in ["footer", "页脚", "page number", "slide number"]):
        return SemanticRole.FOOTER
    if any(kw in name for kw in ["logo", "icon", "picture"]):
        return SemanticRole.LOGO
    if any(kw in name for kw in ["divider", "line", "分隔", "直线", "connector"]):
        return SemanticRole.DIVIDER
    if any(kw in name for kw in ["decor", "accent", "shape", "rectangle", "矩形", "oval", "椭圆"]):
        return SemanticRole.DECORATION

    # Text-based inference for shapes with text
    if shape.has_text_frame and text_len > 0:
        # Short bold text → likely title
        if text_len < 30 and "\n" not in text.strip():
            return SemanticRole.TITLE
        # Bullet list → body
        if "\n" in text and text_len > 50:
            return SemanticRole.BODY
        # Short footer-like text
        if text_len < 15:
            return SemanticRole.FOOTER

    # No text shape → decoration
    if not shape.has_text_frame or not text.strip():
        return SemanticRole.DECORATION

    return SemanticRole.BODY


# ── Text style extraction ────────────────────────────────────────────


def _extract_text_style(text_frame) -> TextStyle:
    """Extract font details from the first text run across all paragraphs."""
    style = TextStyle()

    if not text_frame.paragraphs:
        return style

    # Find first paragraph with runs
    first_run = None
    first_para = None
    for para in text_frame.paragraphs:
        if para.runs:
            first_para = para
            first_run = para.runs[0]
            break

    if first_para is None:
        return style

    # Alignment from the first paragraph with content
    try:
        from pptx.enum.text import PP_ALIGN
        align_map = {
            PP_ALIGN.LEFT: "left",
            PP_ALIGN.CENTER: "center",
            PP_ALIGN.RIGHT: "right",
            PP_ALIGN.JUSTIFY: "justify",
        }
        style.text_alignment = align_map.get(first_para.alignment, "left")
    except Exception:
        pass

    # Line spacing. python-pptx returns either a float (multiple, e.g. 1.15)
    # or a Length (exact spacing, e.g. Centipoints). The spec field is a
    # multiplier, so only accept numeric multiples — a Length here would leak
    # a python-pptx object into the YAML and break yaml.safe_load downstream.
    try:
        ls = first_para.line_spacing
        if isinstance(ls, (int, float)) and not isinstance(ls, bool) and ls > 0:
            style.line_spacing = float(ls)
    except Exception:
        pass

    # Extract from first run
    font = first_run.font

    # Font family
    if font.name:
        style.font_family = font.name

    # Font size
    if font.size:
        style.font_size_pt = font.size / EMU_PER_PT

    # Bold / italic / underline
    style.font_weight = "bold" if font.bold else "normal"
    style.font_italic = bool(font.italic) if font.italic is not None else False
    style.font_underline = bool(font.underline) if font.underline is not None else False

    # Font color — raw rPr XML first (keeps lumMod/tint), then ColorFormat
    try:
        hexv = ""
        try:
            rPr = first_run._r.find(f"{{{_NS_A}}}rPr")
            solid = rPr.find(f"{{{_NS_A}}}solidFill") if rPr is not None else None
            if solid is not None:
                hexv = resolve_color_element(_first_color_el(solid))
        except Exception:
            hexv = ""
        if not hexv:
            hexv = resolve_to_hex(font.color)
        if hexv:
            style.font_color = hexv
    except Exception:
        pass

    return style


# ── Shape style extraction ────────────────────────────────────────────


def _extract_shape_style(shape) -> ShapeStyle:
    """Extract fill, stroke, and geometry details from a shape."""
    style = ShapeStyle()

    spPr = None
    try:
        spPr = shape._element.spPr
    except Exception:
        spPr = None

    # Fill — prefer the raw spPr XML so lumMod/lumOff tints and gradients are
    # reproduced exactly rather than approximated by the base theme colour.
    if spPr is not None:
        try:
            solid = spPr.find(f"{{{_NS_A}}}solidFill")
            if solid is not None:
                color_el = _first_color_el(solid)
                hexv = resolve_color_element(color_el)
                if hexv:
                    style.fill_color = hexv
                    style.fill_type = "solid"
                    style.fill_opacity = color_alpha(color_el)
            else:
                grad = spPr.find(f"{{{_NS_A}}}gradFill")
                stops, ang = extract_gradient(grad)
                if stops:
                    style.fill_type = "gradient"
                    style.fill_gradient = stops
                    style.fill_gradient_angle = ang
                    style.fill_color = stops[0]["color"]
                    if len(stops) > 1:
                        style.fill_color = stops[-1]["color"]
        except Exception:
            pass

    # Fill fallback via python-pptx (for shapes whose spPr has no colour child)
    if not style.fill_color:
        try:
            fill = shape.fill
            from pptx.enum.dml import MSO_FILL_TYPE
            if fill.type == MSO_FILL_TYPE.SOLID:
                hexv = resolve_to_hex(fill.fore_color)
                if hexv:
                    style.fill_color = hexv
                    style.fill_type = "solid"
        except Exception:
            pass

    # Stroke / Line — raw XML first (keeps tints), then python-pptx fallback.
    if spPr is not None:
        try:
            ln = spPr.find(f"{{{_NS_A}}}ln")
            if ln is not None:
                solid = ln.find(f"{{{_NS_A}}}solidFill")
                if solid is not None:
                    color_el = _first_color_el(solid)
                    hexv = resolve_color_element(color_el)
                    if hexv:
                        style.stroke_color = hexv
                        style.stroke_opacity = color_alpha(color_el)
                if ln.get("w"):
                    style.stroke_width_pt = int(ln.get("w")) / EMU_PER_PT
        except Exception:
            pass

    if not style.stroke_color:
        try:
            if hasattr(shape, "line") and shape.line:
                line = shape.line
                from pptx.enum.dml import MSO_FILL_TYPE
                if line.fill.type == MSO_FILL_TYPE.SOLID:
                    hexv = resolve_to_hex(line.fill.fore_color)
                    if hexv:
                        style.stroke_color = hexv
                if line.width and not style.stroke_width_pt:
                    style.stroke_width_pt = line.width / EMU_PER_PT
        except Exception:
            pass

    # Geometry type
    try:
        style.shape_type = _get_geometry_type(shape)
    except Exception:
        pass

    # Corner radius for rounded rects
    try:
        if hasattr(shape, "_element"):
            xml = shape._element.xml
            import re
            m = re.search(r'<a:prstGeom[^>]*prst="([^"]*)"', xml)
            if m:
                style.shape_type = m.group(1)
            # Look for rounded corner radius
            m2 = re.search(r'<a:roundRect[^>]*av="(\d+)"', xml)
            if m2:
                style.corner_radius_pt = int(m2.group(1)) / EMU_PER_PT
    except Exception:
        pass

    return style


def _get_geometry_type(shape) -> str:
    """Get human-readable geometry type name."""
    if not hasattr(shape, "shape_type") or shape.shape_type is None:
        return "rect"

    st = str(shape.shape_type).upper()

    if "RECTANGLE" in st:
        # Check if rounded
        try:
            xml = shape._element.xml
            if "roundRect" in xml:
                return "round_rect"
        except Exception:
            pass
        return "rect"
    if "ROUNDED" in st:
        return "round_rect"
    if "OVAL" in st or "ELLIPSE" in st:
        return "circle"
    if "LINE" in st or "CONNECTOR" in st:
        return "line"
    if "TRIANGLE" in st:
        return "triangle"
    if "ARROW" in st or "CHEVRON" in st:
        return "arrow"
    if "PARALLELOGRAM" in st:
        return "parallelogram"
    if "DIAMOND" in st or "RHOMBUS" in st:
        return "diamond"
    if "HEXAGON" in st:
        return "hexagon"
    if "PENTAGON" in st:
        return "pentagon"
    if "CIRCLE" in st:
        return "circle"
    if "TRAPEZOID" in st:
        return "trapezoid"
    if "PIE" in st:
        return "pie"
    if "CLOUD" in st or "CALLOUT" in st:
        return "callout"
    if "STAR" in st:
        return "star"
    if "WAVE" in st:
        return "wave"
    if "GROUP" in st:
        return "group"
    if "PLACEHOLDER" in st:
        return "placeholder"

    return "rect"


# ── Image extraction ─────────────────────────────────────────────────


def _extract_image_style(picture, spec_dir: Path, idx: int) -> ImageStyle | None:
    """Extract image details and save to assets/."""
    try:
        image = picture.image
        ext = image.content_type.split("/")[-1]
        if ext == "jpeg":
            ext = "jpg"

        assets_dir = spec_dir / "assets"
        assets_dir.mkdir(parents=True, exist_ok=True)

        filename = f"img_{idx:03d}.{ext}"
        filepath = assets_dir / filename
        filepath.write_bytes(image.blob)

        return ImageStyle(
            src=f"assets/{filename}",
            original_width_px=getattr(image, "width", 0) or 0,
            original_height_px=getattr(image, "height", 0) or 0,
            crop_left=getattr(picture, "crop_left", 0.0) or 0.0,
            crop_right=getattr(picture, "crop_right", 0.0) or 0.0,
            crop_top=getattr(picture, "crop_top", 0.0) or 0.0,
            crop_bottom=getattr(picture, "crop_bottom", 0.0) or 0.0,
        )
    except Exception:
        return None


# ── Visual hierarchy summary ──────────────────────────────────────────


def build_hierarchy_summary(elements: list[Element]) -> list[str]:
    """Build a human-readable visual hierarchy description."""
    summary: list[str] = []
    seen_roles: set[str] = set()

    for elem in sorted(elements, key=lambda e: (e.position.y, e.position.x)):
        role = elem.semantic_role.value
        if role in seen_roles:
            continue
        seen_roles.add(role)

        if elem.text_style and elem.text_style.font_family:
            ts = elem.text_style
            summary.append(
                f"{role}: {ts.font_family} {ts.font_size_pt:.0f}pt "
                f"{ts.font_weight} {ts.font_color}"
            )
        elif elem.shape_style and elem.shape_style.fill_color:
            summary.append(
                f"{role}: {elem.shape_style.shape_type} "
                f"fill={elem.shape_style.fill_color}"
            )
        elif elem.image_style:
            summary.append(f"{role}: image")

    return summary
