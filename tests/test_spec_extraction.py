"""The extraction path: a reference .pptx -> specs/<name>/ (spec.yaml + pages/)."""

from __future__ import annotations

import pytest
import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ppt_clone.cli.spec_commands import get_active_spec, list_specs, select_spec
from ppt_clone.spec.enhanced_extractor import SpecExtractor


@pytest.fixture
def reference(tmp_path):
    """A tiny deck with a branded band, a cover title and an agenda page."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    band = cover.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(2))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x7E, 0x69, 0x4F)
    title = cover.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11), Inches(1))
    title.text_frame.text = "示例参考稿封面"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)

    agenda = prs.slides.add_slide(prs.slide_layouts[6])
    agenda.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(5), Inches(0.8)) \
        .text_frame.text = "目录"
    for i in (1, 2, 3):
        agenda.shapes.add_textbox(Inches(0.8), Inches(1.6 + i * 0.6), Inches(5), Inches(0.5)) \
            .text_frame.text = f"{i} 章节"

    path = tmp_path / "reference.pptx"
    prs.save(path)
    return path


def _extract(tmp_path, reference, name="ref"):
    extractor = SpecExtractor()
    spec = extractor.extract(reference, name=name)
    extractor.save(spec)
    return spec


def test_extracts_palette_and_pages(reference, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    spec = _extract(tmp_path, reference)

    assert spec.metadata["name"] == "ref"
    palette = spec.palette.to_dict()
    assert palette and not any(v is None for v in palette.values())
    assert spec.pages
    assert "toc" in spec.page_types_found


def test_writes_the_spec_directory(reference, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _extract(tmp_path, reference)

    spec_dir = tmp_path / "specs" / "ref"
    assert (spec_dir / "spec.yaml").is_file()
    assert (spec_dir / "pages").is_dir()


def test_saved_yaml_reloads(reference, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _extract(tmp_path, reference)

    data = yaml.safe_load((tmp_path / "specs" / "ref" / "spec.yaml").read_text())
    assert data["metadata"]["name"] == "ref"


def test_list_and_select_specs(reference, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _extract(tmp_path, reference)

    assert list_specs() == ["ref"]
    select_spec("ref")
    assert get_active_spec() == "ref"
    with pytest.raises(FileNotFoundError):
        select_spec("no-such-spec")
