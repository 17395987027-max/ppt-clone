"""Image assets: which picture is the ground, and how the rest are used."""

from __future__ import annotations

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from ppt_clone.spec.enhanced_extractor import SpecExtractor


def _deck(tmp_path, build) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(prs, slide, tmp_path)
    return prs


def _save(prs, tmp_path, name="ref.pptx", spec_name="ref"):
    path = tmp_path / name
    prs.save(path)
    return SpecExtractor().extract(path, name=spec_name)


def _png(path, color=(94, 55, 28), size=(800, 450)):
    Image.new("RGB", size, color).save(path)
    return path


def test_full_bleed_picture_becomes_the_background(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    def build(prs, slide, tmp):
        bg = _png(tmp / "bg.png")
        slide.shapes.add_picture(str(bg), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)) \
            .text_frame.text = "封面标题"

    spec = _save(_deck(tmp_path, build), tmp_path, spec_name="bgref")
    page = spec.pages[0]

    assert page.background_type == "image"
    assert page.background_image.startswith("assets/")
    assert any(a.get("usage") == "background" for a in page.assets)


def test_nested_pictures_are_a_mockup_and_its_inset(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    def build(prs, slide, tmp):
        frame = _png(tmp / "frame.png", (30, 30, 30), (600, 400))
        inner = _png(tmp / "inner.png", (200, 120, 60), (600, 400))
        slide.shapes.add_picture(str(frame), Inches(2), Inches(1),
                                 width=Inches(6), height=Inches(4))
        slide.shapes.add_picture(str(inner), Inches(3.5), Inches(2),
                                 width=Inches(3), height=Inches(2))

    spec = _save(_deck(tmp_path, build), tmp_path, spec_name="mock")
    assets = spec.pages[0].assets

    frame = next(a for a in assets if a.get("usage") == "mockup_frame")
    inset = next(a for a in assets if a.get("usage") == "inset")
    assert frame["contains"] == inset["src"]
    assert inset["inside"] == frame["src"]


def test_a_reused_picture_is_saved_once(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    def build(prs, slide, tmp):
        logo = _png(tmp / "logo.png", (180, 140, 90), (200, 200))
        for x in (Inches(1), Inches(4)):
            slide.shapes.add_picture(str(logo), x, Inches(1),
                                     width=Inches(1), height=Inches(1))

    spec = _save(_deck(tmp_path, build), tmp_path, spec_name="dup")
    assets = spec.pages[0].assets

    assert len(assets) == 1
    # Element-level images are saved separately as img_*; only one copy of the
    # reused picture itself should exist.
    saved = [p for p in (tmp_path / "specs" / "dup" / "assets").glob("*")
             if not p.name.startswith("img_")]
    assert len(saved) == 1
