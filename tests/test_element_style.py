"""Unit tests for raw-XML style extraction.

Covers the fidelity fixes that python-pptx's ColorFormat cannot express:
luminance tints (lumMod/lumOff), transparency (a:alpha), and per-stop
gradients. These are read straight from the DrawingML.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from ppt_clone.spec.element_extractor import (
    _extract_shape_style,
    color_alpha,
    resolve_color_element,
    set_theme_colors,
)

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _color_el(xml: str):
    return etree.fromstring(xml)


def test_color_alpha_percent():
    el = _color_el(f'<a:solidFill xmlns:a="{NS_A}">'
                   '<a:srgbClr val="2A2A2A"><a:alpha val="20000"/></a:srgbClr>'
                   '</a:solidFill>')
    srgb = el.find(f"{{{NS_A}}}srgbClr")
    assert color_alpha(srgb) == 0.2


def test_color_alpha_default_opaque():
    el = _color_el(f'<a:srgbClr xmlns:a="{NS_A}" val="FFFFFF"/>')
    assert color_alpha(el) == 1.0


def test_resolve_tint_lummod_lumoff():
    set_theme_colors({"accent1": "336699"})
    el = _color_el(f'<a:schemeClr xmlns:a="{NS_A}" val="accent1">'
                   '<a:lumMod val="40000"/><a:lumOff val="60000"/></a:schemeClr>')
    hexv = resolve_color_element(el)
    # Lightened accent1 must not equal the base colour.
    assert hexv.startswith("#") and hexv.upper() != "#336699"


def test_shape_style_captures_fill_opacity():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shp = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(5))
    spPr = shp._element.spPr
    for tag in ("solidFill", "gradFill", "noFill"):
        node = spPr.find(f"{{{NS_A}}}{tag}")
        if node is not None:
            spPr.remove(node)
    xml = (f'<a:solidFill xmlns:a="{NS_A}">'
           '<a:srgbClr val="2A2A2A"><a:alpha val="20000"/></a:srgbClr>'
           '</a:solidFill>')
    spPr.insert(1, etree.fromstring(xml))

    style = _extract_shape_style(shp)
    assert style.fill_color == "#2A2A2A"
    assert style.fill_type == "solid"
    assert abs(style.fill_opacity - 0.2) < 1e-6
    assert style.to_dict()["fill_opacity"] == 0.2


def test_shape_style_gradient_stops():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shp = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(4), Inches(2))
    spPr = shp._element.spPr
    for tag in ("solidFill", "gradFill", "noFill"):
        node = spPr.find(f"{{{NS_A}}}{tag}")
        if node is not None:
            spPr.remove(node)
    xml = (f'<a:gradFill xmlns:a="{NS_A}"><a:gsLst>'
           '<a:gs pos="0"><a:srgbClr val="FFFFFF"/></a:gs>'
           '<a:gs pos="100000"><a:srgbClr val="000000"/></a:gs>'
           '</a:gsLst><a:lin ang="5400000"/></a:gradFill>')
    spPr.insert(1, etree.fromstring(xml))

    style = _extract_shape_style(shp)
    assert style.fill_type == "gradient"
    assert len(style.fill_gradient) == 2
    assert style.fill_gradient[0]["color"] == "#FFFFFF"
    assert style.fill_gradient_angle == 90.0
